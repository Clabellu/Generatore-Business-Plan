import os
from flask import Flask, render_template, request, jsonify, Response, send_file, flash, redirect, url_for
from flask_cors import CORS
from flask_login import login_required, current_user
from anthropic import Anthropic # Assicurati che sia importato
from dotenv import load_dotenv
from weasyprint import HTML, CSS
from docx import Document
import io
import re
from bs4 import BeautifulSoup
import html

# Import database e auth
from database import init_database, create_tables, db
from auth import init_auth
from models import User, Order
from credits import (
    aggiungi_crediti,
    consuma_credito,
    verifica_crediti,
    get_pacchetti_crediti,
    get_pacchetto_by_id
)

load_dotenv()

app = Flask(__name__)
CORS(app)

# Configurazione SECRET_KEY per le sessioni
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'dev-secret-key-change-in-production')

# Inizializza database
init_database(app)

# Crea tabelle se non esistono
create_tables(app)

# Inizializza sistema autenticazione
init_auth(app)

# Inizializzazione del client Anthropic (come l'avevamo prima)
try:
    api_key = os.getenv("ANTHROPIC_API_KEY")
    if not api_key:
        raise ValueError("La variabile d'ambiente ANTHROPIC_API_KEY non è stata impostata.")
    
    anthropic_client = Anthropic(api_key=api_key) # Rinominato per chiarezza
    print("Client Anthropic inizializzato con successo.")
except Exception as e:
    print(f"Errore fatale durante l'inizializzazione del client Anthropic: {e}")
    anthropic_client = None




def convert_text_to_html(text):
    """
    Converte un testo con formattazione simile a Markdown in HTML pulito e ben strutturato.
    
    Supporta:
    - Titoli (h1-h6): #, ##, ###, ####, #####, ######
    - Grassetto: **testo** o __testo__
    - Corsivo: *testo* o _testo_
    - Codice inline: `codice`
    - Link: [testo](url)
    - Liste ordinate e non ordinate (anche annidate)
    - Tabelle con supporto allineamento
    - Blocchi di codice: ```linguaggio
    - Citazioni: > testo
    - Linee orizzontali: --- o ***
    """
    if not text:
        return ""

    # Escape dei caratteri HTML speciali per sicurezza
    text = html.escape(text)
    
    # Pre-elaborazione per elementi inline
    text = _process_inline_elements(text)
    
    lines = text.split('\n')
    html_output = []
    state = "NORMAL"
    code_lang = ""  # Per blocchi di codice
    table_alignments = []  # Per memorizzare allineamenti tabelle
    list_stack = []  # Stack per liste annidate
    
    i = 0
    while i < len(lines):
        line = lines[i]
        line_stripped = line.strip()
        original_line = line  # Mantieni indentazione originale per liste annidate
        
        # --- GESTIONE USCITA DA STATI PRECEDENTI ---
        if state == "IN_CODE_BLOCK" and line_stripped == "```":
            html_output.append('</code></pre>')
            state = "NORMAL"
            i += 1
            continue
            
        if state == "IN_CODE_BLOCK":
            html_output.append(html.escape(line))
            i += 1
            continue
            
        if state == "IN_TABLE" and not line_stripped.startswith('|'):
            html_output.append('</tbody></table>')
            state = "NORMAL"
            table_alignments = []
            
        # Gestione chiusura liste annidate
        if state.startswith("IN_LIST"):
            current_indent = len(original_line) - len(original_line.lstrip())
            if not _is_list_item(line_stripped):
                # Chiudi tutte le liste aperte
                while list_stack:
                    list_type = list_stack.pop()
                    html_output.append(f'</{list_type}>')
                state = "NORMAL"
            else:
                # Gestisci annidamento
                _handle_list_nesting(html_output, list_stack, current_indent, line_stripped)
                i += 1
                continue
        
        # --- ANALISI DELLA RIGA CORRENTE ---
        
        # Gestione BLOCCHI DI CODICE
        if line_stripped.startswith('```'):
            if state != "IN_CODE_BLOCK":
                state = "IN_CODE_BLOCK"
                code_lang = line_stripped[3:].strip()
                lang_attr = f' class="language-{code_lang}"' if code_lang else ''
                html_output.append(f'<pre><code{lang_attr}>')
            i += 1
            continue
            
        # Gestione LINEE ORIZZONTALI
        elif line_stripped in ['---', '***', '___']:
            html_output.append('<hr>')
            
        # Gestione CITAZIONI
        elif line_stripped.startswith('> '):
            quote_text = line_stripped[2:]
            html_output.append(f'<blockquote><p>{quote_text}</p></blockquote>')
            
        # Gestione TITOLI (h1-h6)
        elif line_stripped.startswith('#'):
            level = 0
            for char in line_stripped:
                if char == '#':
                    level += 1
                else:
                    break
            if 1 <= level <= 6 and line_stripped[level:level+1] == ' ':
                title_text = line_stripped[level+1:].strip()
                html_output.append(f'<h{level}>{title_text}</h{level}>')
            else:
                html_output.append(f'<p>{line_stripped}</p>')
                
        # Gestione TABELLE
        elif line_stripped.startswith('|') and line_stripped.endswith('|'):
            if state != "IN_TABLE":
                # Prima riga: intestazione
                state = "IN_TABLE"
                parts = [p.strip() for p in line_stripped.strip('|').split('|')]
                html_output.append('<table class="financial-table">')
                html_output.append('<thead><tr>')
                for part in parts:
                    html_output.append(f'<th>{part}</th>')
                html_output.append('</tr></thead>')
                
                # Controlla la prossima riga per allineamenti
                if i + 1 < len(lines):
                    next_line = lines[i + 1].strip()
                    if next_line.startswith('|') and next_line.endswith('|'):
                        align_parts = [p.strip() for p in next_line.strip('|').split('|')]
                        if all(re.match(r'^:?-+:?$', part) for part in align_parts):
                            table_alignments = []
                            for part in align_parts:
                                if part.startswith(':') and part.endswith(':'):
                                    table_alignments.append('center')
                                elif part.endswith(':'):
                                    table_alignments.append('right')
                                else:
                                    table_alignments.append('left')
                            i += 1  # Salta la riga di allineamento
                
                html_output.append('<tbody>')
            else:
                # Riga di dati
                parts = [p.strip() for p in line_stripped.strip('|').split('|')]
                html_output.append('<tr>')
                for j, part in enumerate(parts):
                    align = ''
                    if j < len(table_alignments) and table_alignments[j] != 'left':
                        align = f' style="text-align: {table_alignments[j]}"'
                    html_output.append(f'<td{align}>{part}</td>')
                html_output.append('</tr>')
                
        # Gestione LISTE
        elif _is_list_item(line_stripped):
            if not state.startswith("IN_LIST"):
                state = "IN_LIST"
                list_stack = []
            
            current_indent = len(original_line) - len(original_line.lstrip())
            _handle_list_nesting(html_output, list_stack, current_indent, line_stripped)
            
        # Gestione PARAGRAFI
        elif line_stripped:
            html_output.append(f'<p>{line_stripped}</p>')
            
        i += 1
    
    # --- CHIUSURA FINALE ---
    if state == "IN_CODE_BLOCK":
        html_output.append('</code></pre>')
    elif state == "IN_TABLE":
        html_output.append('</tbody></table>')
    elif state.startswith("IN_LIST"):
        while list_stack:
            list_type = list_stack.pop()
            html_output.append(f'</{list_type}>')
    
    return '\n'.join(html_output)


def _process_inline_elements(text):
    """Elabora elementi inline come grassetto, corsivo, codice, link."""
    # Codice inline (prima di grassetto/corsivo per evitare conflitti)
    text = re.sub(r'`([^`]+)`', r'<code>\1</code>', text)
    
    # Link [testo](url)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'<a href="\2">\1</a>', text)
    
    # Grassetto **testo** o __testo__
    text = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', text)
    text = re.sub(r'__(.*?)__', r'<strong>\1</strong>', text)
    
    # Corsivo *testo* o _testo_ (dopo grassetto per evitare conflitti)
    text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'<em>\1</em>', text)
    text = re.sub(r'(?<!_)_([^_]+)_(?!_)', r'<em>\1</em>', text)
    
    return text


def _is_list_item(line):
    """Verifica se una riga è un elemento di lista."""
    return (line.startswith(('-', '* ', '+ ')) or 
            re.match(r'^\d+[\.)]\s', line))


def _handle_list_nesting(html_output, list_stack, current_indent, line_stripped):
    """Gestisce l'annidamento delle liste."""
    # Determina il tipo di lista
    if re.match(r'^\d+[\.)]\s', line_stripped):
        current_list_type = 'ol'
        item_text = re.sub(r'^\d+[\.)]\s', '', line_stripped)
    else:
        current_list_type = 'ul'
        item_text = line_stripped[2:]  # Rimuovi '- ', '* ', o '+ '
    
    # Gestisci stack delle liste per annidamento
    target_depth = current_indent // 2  # Assumiamo 2 spazi per livello
    
    # Chiudi liste se necessario
    while len(list_stack) > target_depth:
        list_type = list_stack.pop()
        html_output.append(f'</{list_type}>')
    
    # Apri nuove liste se necessario
    while len(list_stack) < target_depth:
        list_stack.append(current_list_type)
        html_output.append(f'<{current_list_type}>')
    
    # Se non ci sono liste aperte o il tipo è diverso
    if not list_stack or list_stack[-1] != current_list_type:
        if list_stack:
            # Sostituisci l'ultimo tipo
            old_type = list_stack.pop()
            html_output.append(f'</{old_type}>')
        list_stack.append(current_list_type)
        html_output.append(f'<{current_list_type}>')
    
    html_output.append(f'<li>{item_text}</li>')

def calcola_e_formatta_proiezioni(dati_finanziari, valuta_simbolo="€"):
    """
    Calcola e formatta le proiezioni finanziarie a 5 anni basandosi sui dati del form.
    Versione corretta con debug completo.
    """
    try:
        print("=== DEBUG CALCOLO PROIEZIONI ===")
        print(f"Dati ricevuti: {dati_finanziari}")
        
        # Estrai dati base
        vendite_anno1 = float(dati_finanziari.get('venditeAnno1', 0))
        crescita_percent = float(dati_finanziari.get('crescitaFatturatoAnnuale', 0))
        
        print(f"Vendite Anno 1: {vendite_anno1}")
        print(f"Crescita %: {crescita_percent}")
        
        # Estrai dati costi dalla tabella
        costi_table = dati_finanziari.get('costiTable', [])
        print(f"Costi table ricevuti: {costi_table}")
        
        # Mappa percentuali per ID
        percentuali = {}
        for costo in costi_table:
            id_costo = costo.get('id', '')
            perc_costo = float(costo.get('percentuale', 0))
            percentuali[id_costo] = perc_costo
            print(f"  {id_costo}: {perc_costo}%")
        
        print(f"Percentuali finali: {percentuali}")
        
        # Se non ci sono dati, usa valori di esempio per debug
        if vendite_anno1 == 0:
            print("ATTENZIONE: Vendite = 0, usando valori di esempio per debug")
            vendite_anno1 = 100000
            crescita_percent = 10
            percentuali = {
                'costoDelVenduto': 40,
                'salariEBenefici': 6,
                'marketing': 5,
                'affitto': 0,
                'generaleEAmministrazione': 1,
                'ammortamento': 2,
                'costiAccessori': 0,
                'altreSpese': 1,
                'interessiPassivi': 0,
                'importoTasse': 20
            }
        
        # Inizializza struttura dati per 5 anni
        anni = 5
        proiezioni = {}
        
        # Calcola proiezioni anno per anno
        vendite_correnti = vendite_anno1
        
        # Voci di costo con mapping ID -> Nome
        voci_mapping = {
            'costoDelVenduto': 'Costo del venduto',
            'salariEBenefici': 'Salari e benefici',
            'marketing': 'Marketing',
            'affitto': 'Affitto',
            'generaleEAmministrazione': 'Generale e amministrazione',
            'ammortamento': 'Ammortamento',
            'costiAccessori': 'Costi accessori',
            'altreSpese': 'Altre spese',
            'interessiPassivi': 'Interessi passivi'
        }
        
        # Inizializza tutte le voci
        proiezioni['Vendite'] = []
        for nome_voce in voci_mapping.values():
            proiezioni[nome_voce] = []
        proiezioni['Importo tasse (% EBT)'] = []
        
        for anno in range(anni):
            print(f"\n--- ANNO {anno + 1} ---")
            print(f"Vendite: {vendite_correnti:,.2f}")
            
            # Vendite per questo anno
            proiezioni['Vendite'].append(vendite_correnti)
            
            # Calcola ogni voce di costo
            costi_operativi_totali = 0
            interessi_passivi = 0
            
            for id_costo, nome_costo in voci_mapping.items():
                percentuale = percentuali.get(id_costo, 0)
                costo_anno = vendite_correnti * (percentuale / 100)
                proiezioni[nome_costo].append(costo_anno)
                
                print(f"  {nome_costo}: {percentuale}% = {costo_anno:,.2f}")
                
                if id_costo == 'interessiPassivi':
                    interessi_passivi = costo_anno
                else:
                    costi_operativi_totali += costo_anno
            
            # Calcola EBT (Utili prima delle tasse)
            ebt = vendite_correnti - costi_operativi_totali - interessi_passivi
            print(f"  EBT: {ebt:,.2f}")
            
            # Tasse (percentuale su EBT, solo se EBT > 0)
            percentuale_tasse = percentuali.get('importoTasse', 20)
            tasse = max(0, ebt * (percentuale_tasse / 100))
            proiezioni['Importo tasse (% EBT)'].append(tasse)
            print(f"  Tasse ({percentuale_tasse}%): {tasse:,.2f}")
            
            # Prossimo anno
            vendite_correnti *= (1 + crescita_percent / 100)
        
        print(f"\nProiezioni finali: {proiezioni}")
        
        # Costruisci la tabella Markdown
        tabella_risultato = _costruisci_tabella_semplice_robusta(proiezioni, valuta_simbolo)
        
        return tabella_risultato
    
    
    except Exception as e:
        print(f"ERRORE nel calcolo proiezioni: {e}")
        import traceback
        traceback.print_exc()
        return _tabella_errore(valuta_simbolo)


def _costruisci_tabella_semplice_robusta(proiezioni, valuta_simbolo="€"):
    """
    Costruisce una tabella finanziaria completa con tutti i calcoli intermedi
    inclusi Margine Lordo, Spese Totali, EBIT, EBT, etc.
    """
    
    print("=== COSTRUZIONE TABELLA FINANZIARIA COMPLETA ===")
    
    # Header della tabella
    table_md = "| Voce Finanziaria | Anno 1 | Anno 2 | Anno 3 | Anno 4 | Anno 5 |\n"
    table_md += "|:---|---:|---:|---:|---:|---:|\n"
    
    try:
        # === SEZIONE 1: ENTRATE ===
        if 'Vendite' in proiezioni:
            vendite = proiezioni['Vendite']
            vendite_formattate = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in vendite]
            table_md += "| **Entrate** | **" + "** | **".join(vendite_formattate) + "** |\n"
        
        # === SEZIONE 2: COSTO DEL VENDUTO ===
        if 'Costo del venduto' in proiezioni:
            costo_venduto = proiezioni['Costo del venduto']
            costo_formattato = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in costo_venduto]
            table_md += "| Costo del Venduto | " + " | ".join(costo_formattato) + " |\n"
        
        # === CALCOLO E SEZIONE 3: MARGINE LORDO ===
        if 'Vendite' in proiezioni and 'Costo del venduto' in proiezioni:
            margine_lordo = []
            for i in range(5):
                margine = proiezioni['Vendite'][i] - proiezioni['Costo del venduto'][i]
                margine_lordo.append(margine)
            
            margine_formattato = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in margine_lordo]
            table_md += "| **Margine Lordo** | **" + "** | **".join(margine_formattato) + "** |\n"
        
        # Riga vuota per separazione
        table_md += "| | | | | | |\n"
        
        # === SEZIONE 4: SPESE OPERATIVE ===
        table_md += "| **Spese operative** | | | | | |\n"
        
        # Lista delle spese operative in ordine
        spese_operative = [
            'Salari e benefici',
            'Marketing',
            'Affitto',
            'Generale e amministrazione',
            'Ammortamento',
            'Costi accessori',  # Utenze nel tuo screenshot
            'Altre spese'
        ]
        
        # Calcola spese totali mentre aggiungi le singole voci
        spese_totali = [0, 0, 0, 0, 0]  # Array per 5 anni
        
        for spesa in spese_operative:
            if spesa in proiezioni:
                valori_spesa = proiezioni[spesa]
                # Aggiungi ai totali
                for i in range(5):
                    spese_totali[i] += valori_spesa[i]
                
                # Formatta e aggiungi alla tabella
                valori_formattati = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in valori_spesa]
                # Indenta le spese operative
                nome_indentato = f"    {spesa}"
                table_md += f"| {nome_indentato} | " + " | ".join(valori_formattati) + " |\n"
        
        # === SEZIONE 5: SPESE TOTALI ===
        spese_totali_formattate = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in spese_totali]
        table_md += "| **Spese totali** | **" + "** | **".join(spese_totali_formattate) + "** |\n"
        
        # Riga vuota per separazione
        table_md += "| | | | | | |\n"
        
        # === CALCOLO E SEZIONE 6: UTILI PRIMA DEGLI INTERESSI E DELLE TASSE (EBIT) ===
        if margine_lordo:
            ebit = []
            for i in range(5):
                utile = margine_lordo[i] - spese_totali[i]
                ebit.append(utile)
            
            ebit_formattato = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in ebit]
            table_md += "| **Utili prima degli interessi e delle tasse** | **" + "** | **".join(ebit_formattato) + "** |\n"
        
        # === SEZIONE 7: INTERESSI PASSIVI ===
        if 'Interessi passivi' in proiezioni:
            interessi = proiezioni['Interessi passivi']
            interessi_formattati = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in interessi]
            table_md += "| Spese per interessi | " + " | ".join(interessi_formattati) + " |\n"
        else:
            # Se non ci sono interessi, usa zero
            interessi = [0, 0, 0, 0, 0]
            interessi_formattati = [f"0 {valuta_simbolo}" for _ in range(5)]
            table_md += "| Spese per interessi | " + " | ".join(interessi_formattati) + " |\n"
        
        # === CALCOLO E SEZIONE 8: UTILI PRIMA DELLE TASSE (EBT) ===
        if ebit:
            ebt = []
            for i in range(5):
                utile_pre_tasse = ebit[i] - interessi[i]
                ebt.append(utile_pre_tasse)
            
            ebt_formattato = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in ebt]
            table_md += "| **Utili prima delle tasse** | **" + "** | **".join(ebt_formattato) + "** |\n"
        
        # === SEZIONE 9: TASSE ===
        if 'Importo tasse (% EBT)' in proiezioni:
            tasse = proiezioni['Importo tasse (% EBT)']
            tasse_formattate = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in tasse]
            table_md += "| Imposte sul reddito | " + " | ".join(tasse_formattate) + " |\n"
        
        # Riga vuota per separazione finale
        table_md += "| | | | | | |\n"
        
        # === CALCOLO E SEZIONE 10: REDDITO NETTO FINALE ===
        if ebt and 'Importo tasse (% EBT)' in proiezioni:
            reddito_netto = []
            for i in range(5):
                netto = ebt[i] - proiezioni['Importo tasse (% EBT)'][i]
                reddito_netto.append(netto)
            
            netto_formattato = [f"{val:,.0f} {valuta_simbolo}".replace(",", ".") for val in reddito_netto]
            table_md += "| **Reddito netto** | **" + "** | **".join(netto_formattato) + "** |\n"
        
        print("✅ Tabella finanziaria completa generata con successo")
        print(f"Lunghezza finale: {len(table_md)} caratteri")
        
        return table_md
        
    except Exception as e:
        print(f"❌ Errore nella costruzione della tabella completa: {e}")
        import traceback
        traceback.print_exc()
        return _tabella_errore(valuta_simbolo)

def _tabella_errore(valuta_simbolo="€"):
    """Restituisce una tabella di errore."""
    return f"""
| Voce Finanziaria | Anno 1 | Anno 2 | Anno 3 | Anno 4 | Anno 5 |
|:---|---:|---:|---:|---:|---:|
| **Errore nel calcolo** | 0 {valuta_simbolo} | 0 {valuta_simbolo} | 0 {valuta_simbolo} | 0 {valuta_simbolo} | 0 {valuta_simbolo} |

*Si è verificato un errore nel calcolo delle proiezioni finanziarie. Verificare i dati inseriti e controllare i log del server.*
"""



# FUNZIONE STANDALONE (FUORI DALLA CLASSE)
def ottieni_istruzioni_per_sezione(nome_sezione, tabella_finanziaria_md="", tipo_bp="completo"):
    """
    Restituisce prompt specifici per ogni sezione del business plan.
    IMPORTANTE: Solo testo narrativo e discorsivo. NO elenchi numerati o puntati.
    Supporta due versioni: 'completo' (dettagliato) e 'breve' (sintetico, max 5 pagine).
    """
    # Prefisso per versione breve
    prefisso_breve = """
VERSIONE SINTETICA: Questo business plan deve essere CONCISO e SINTETICO, limitato a un massimo di 5 pagine totali.
La tua sezione deve essere BREVE ma COMPLETA, limitata a 500-700 parole massimo.
Vai dritto al punto, mantenendo solo le informazioni essenziali.
"""

    istruzioni_specifiche = {
        "Riassunto Esecutivo": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Ogni concetto deve essere espresso attraverso paragrafi fluidi e ben collegati, come se stessi scrivendo un documento aziendale professionale in prosa.

Basandoti esclusivamente sulle informazioni fornite dall'utente nei form, scrivi un Riassunto Esecutivo professionale e convincente che catturi immediatamente l'attenzione degli investitori bancari. Non inventare dettagli non forniti, ma presenta in modo persuasivo e narrativo le informazioni già raccolte.

Inizia presentando l'azienda e la sua missione, utilizzando la descrizione del business fornita per spiegare chiaramente quale problema risolve nel mercato e perché la soluzione è necessaria. Continua descrivendo il prodotto o servizio attraverso una narrazione che evidenzi il valore concreto offerto, basandoti sui benefici specifici menzionati dall'utente.

Sviluppa poi una descrizione coinvolgente del potenziale di mercato, trasformando le informazioni sul target fornite in una storia convincente sulle opportunità di business. Presenta i vantaggi competitivi come elementi distintivi che posizionano l'azienda in modo unico, spiegando attraverso un racconto fluido perché i clienti sceglieranno questa soluzione.

Introduci il team attraverso una presentazione narrativa delle competenze e esperienze, mostrando come questi professionisti siano la chiave del successo futuro. Concludi con una sintesi finanziaria che trasformi i numeri in una storia di crescita sostenibile e opportunità di investimento attraente.

Mantieni sempre un tono professionale ma accessibile, costruendo ogni paragrafo sul precedente per creare un flusso logico e persuasivo che dimostri la solidità dell'opportunità di business.
        """,
        
        "Analisi della Situazione": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Sviluppa ogni punto attraverso paragrafi descrittivi e ben collegati tra loro.

Sviluppa un'analisi dettagliata e narrativa della situazione di mercato, basandoti rigorosamente sulle informazioni fornite dall'utente. Non aggiungere dati di mercato non specificati, ma costruisci una storia professionale e coinvolgente del contesto in cui opera l'azienda.

Apri con una descrizione approfondita del settore, utilizzando le informazioni fornite dall'utente per dipingere un quadro chiaro e dettagliato del mercato di riferimento. Racconta la storia di questo settore, spiegando le dinamiche che lo caratterizzano e perché rappresenta un'opportunità interessante per l'azienda.

Prosegui descrivendo il panorama competitivo attraverso una narrazione che evidenzi come l'azienda si posiziona rispetto ai concorrenti esistenti. Utilizza i punti di differenziazione specificati per costruire un racconto convincente su ciò che rende unica questa proposta di business nel contesto competitivo attuale.

Approfondisci le tendenze di mercato menzionate dall'utente, spiegando attraverso una descrizione fluida come queste evoluzioni creano opportunità specifiche per l'azienda. Presenta le sfide del settore come parte di una storia più ampia, mostrando come l'azienda si prepara ad affrontarle con le strategie descritte.

Sviluppa l'analisi SWOT trasformando ogni elemento in una narrazione dettagliata. Per i punti di forza, racconta come questi vantaggi si traducono in opportunità concrete di successo. Per le debolezze, presenta una descrizione onesta ma costruttiva di come l'azienda sta lavorando per superarle. Per le opportunità, dipingi un quadro delle possibilità future basandoti sui trend identificati. Per le minacce, descrivi i rischi in modo professionale e le strategie per affrontarli.

Concludi con una riflessione narrativa su perché il timing attuale è appropriato per l'azienda, collegando tutti gli elementi analizzati in una visione coerente del contesto di business e delle prospettive future.
        """,
        
        "Marketing": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Presenta la strategia attraverso paragrafi fluidi e descrittivi.

Presenta la strategia di marketing dell'azienda attraverso una narrazione professionale e dettagliata, basandoti esclusivamente sulle informazioni fornite dall'utente. Elabora le informazioni disponibili costruendo una storia coerente e convincente della strategia commerciale, senza inventare tattiche o canali non specificati.

Inizia raccontando chi sono i clienti target dell'azienda, utilizzando la descrizione fornita per creare un ritratto dettagliato dei segmenti di mercato che l'azienda intende servire. Descrivi questi clienti come protagonisti di una storia, spiegando le loro esigenze, le loro sfide quotidiane e come l'azienda risponde perfettamente a queste necessità attraverso la proposta di valore indicata.

Continua sviluppando la strategia di posizionamento attraverso una narrazione che mostri come l'azienda si distingue nel mercato. Basandoti sulla differenziazione competitiva descritta dall'utente, racconta la storia di come questa azienda ha trovato il suo spazio unico nel panorama competitivo. Spiega come la strategia di pricing si integra naturalmente in questo posizionamento, creando un'offerta che risuona con il valore percepito dal mercato target.

Prosegui descrivendo come l'azienda raggiunge i suoi clienti, presentando i canali di marketing e distribuzione specificati attraverso una narrazione che mostri come questi si combinano strategicamente per creare una presenza efficace sul mercato. Racconta come le diverse tattiche di acquisizione clienti lavorano insieme per supportare gli obiettivi di crescita forniti, creando un ecosistema di marketing coerente e mirato.

Approfondisci la gestione del budget di marketing indicato, spiegando attraverso una descrizione dettagliata come le risorse vengono allocate strategicamente tra le diverse attività. Presenta le metriche di performance menzionate dall'utente come strumenti di navigazione che guidano l'azienda verso il successo, spiegando come vengono utilizzate per ottimizzare continuamente l'approccio al mercato.

Concludi con una riflessione su come l'intera strategia di marketing si allinea perfettamente con il modello di business descritto, creando un approccio integrato che massimizza le opportunità nel mercato di riferimento.
        """,
        
        "Operazioni": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Descrivi il piano operativo attraverso paragrafi fluidi e dettagliati.

Descrivi il piano operativo dell'azienda attraverso una narrazione professionale e dettagliata, utilizzando esclusivamente le informazioni fornite dall'utente. Elabora le informazioni disponibili costruendo una storia coerente di come l'azienda funziona quotidianamente, senza aggiungere dettagli operativi non specificati.

Apri presentando la struttura organizzativa dell'azienda come il fondamento su cui si costruisce il successo operativo. Utilizza la descrizione fornita dall'utente per raccontare come i ruoli e le responsabilità si integrano per supportare l'esecuzione del business plan. Descrivi i processi operativi chiave attraverso una narrazione che mostri come questi garantiscono la qualità del prodotto o servizio e la soddisfazione del cliente.

Continua esplorando l'infrastruttura tecnologica e gli strumenti utilizzati dall'azienda, basandoti sulle informazioni fornite per spiegare come la tecnologia supporta operativamente la crescita pianificata. Presenta i requisiti di risorse umane indicati come parte di una strategia più ampia, collegandoli naturalmente alle fasi di sviluppo del business descritte.

Approfondisci le relazioni con fornitori e partner operativi menzionati dall'utente, raccontando come queste collaborazioni contribuiscono all'efficienza complessiva dell'azienda. Descrivi i controlli di qualità e gli standard operativi specificati come elementi che garantiscono l'eccellenza nell'erogazione del servizio, integrando queste pratiche nella narrazione più ampia dell'operatività aziendale.

Sviluppa il tema della scalabilità operativa basandoti sulle informazioni fornite riguardo alla crescita prevista, spiegando attraverso una descrizione dettagliata come l'azienda si prepara ad espandere le proprie operazioni mantenendo gli standard qualitativi. Presenta le metriche operative indicate dall'utente come strumenti di controllo che permettono un monitoraggio continuo e l'ottimizzazione dell'efficienza aziendale.

Concludi collegando tutti gli elementi operativi in una visione coerente che mostri come l'approccio operativo supporta concretamente gli obiettivi di business indicati, dimostrando la solidità della struttura operativa dell'azienda.
        """,
        
        "Gestione": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Presenta il team attraverso paragrafi descrittivi e coinvolgenti.

Presenta il team di gestione dell'azienda attraverso una narrazione professionale e coinvolgente, basandoti rigorosamente sulle informazioni fornite dall'utente. Non aggiungere competenze o esperienze non specificate, ma elabora professionalmente i profili disponibili costruendo una storia convincente del capitale umano dell'azienda.

Inizia raccontando la storia dei fondatori e dei leader dell'azienda, utilizzando le informazioni fornite per presentare le loro competenze, esperienze e ruoli in modo narrativo e coinvolgente. Descrivi come il background professionale di ciascun membro del team li qualifica perfettamente per guidare l'azienda verso gli obiettivi specificati nel business plan, creando un ritratto umano e professionale che ispiri fiducia.

Continua descrivendo la struttura di governance attraverso una narrazione che spieghi come le responsabilità sono distribuite e come vengono prese le decisioni strategiche nell'azienda. Presenta questo sistema come un meccanismo ben oliato che garantisce efficienza e trasparenza nella gestione aziendale. Approfondisci il ruolo dei consiglieri o advisory board menzionati, raccontando come il loro contributo arricchisce e supporta lo sviluppo strategico dell'azienda.

Sviluppa il tema della crescita del team basandoti sulle informazioni sulle assunzioni future indicate dall'utente. Racconta come i nuovi ruoli chiave verranno integrati secondo le tempistiche specificate, spiegando attraverso una descrizione dettagliata come questi contribuiranno alla scalabilità dell'organizzazione e al raggiungimento degli obiettivi di crescita.

Approfondisci la cultura aziendale e i valori utilizzando le informazioni fornite, spiegando come il team intende costruire un ambiente di lavoro che supporti attivamente la mission dell'azienda. Descrivi gli incentivi o i piani di retention specificati come elementi che dimostrano l'attenzione dell'azienda verso il proprio capitale umano e la volontà di mantenere e motivare il talento chiave.

Concludi con una riflessione su come la composizione e le competenze del team rendano credibile e realizzabile l'esecuzione del piano di business proposto, presentando il capitale umano come uno degli asset più importanti per il successo futuro dell'azienda.
        """,
        
        "Strategia di Crescita": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Sviluppa la strategia attraverso paragrafi fluidi e ben collegati.

Presenta la strategia di crescita a lungo termine dell'azienda attraverso una narrazione professionale e visionaria, basandoti esclusivamente sulle informazioni fornite dall'utente. Elabora le informazioni disponibili costruendo una storia convincente del futuro dell'azienda, senza speculare su opportunità non menzionate.

Apri con la vision a lungo termine descritta dall'utente, trasformandola in una narrazione coinvolgente che presenti la direzione strategica dell'azienda nei prossimi anni. Racconta questa vision come una destinazione chiara e raggiungibile, spiegando come gli obiettivi di crescita specificati si collegano naturalmente alle capacità attuali dell'azienda e alle opportunità di mercato identificate.

Continua descrivendo le fasi di crescita indicate dall'utente attraverso una narrazione che mostri come ogni milestone rappresenta un passo logico verso il raggiungimento della vision complessiva. Presenta questo percorso come una strategia ben ponderata, spiegando come ogni fase costruisce sulla precedente per creare un momentum sostenibile di sviluppo aziendale.

Approfondisci i piani di espansione geografica o di mercato specificati, raccontando come questi si integrano perfettamente con le risorse e le competenze disponibili. Descrivi questa espansione come un'evoluzione naturale del business attuale, mostrando come l'azienda sia preparata ad affrontare le sfide e cogliere le opportunità di mercati più ampi.

Sviluppa le strategie di diversificazione o sviluppo prodotto menzionate dall'utente, spiegando attraverso una descrizione dettagliata come queste opportunità emergono organicamente dal business attuale. Presenta le informazioni sulle partnership strategiche o acquisizioni indicate come elementi che accelereranno la crescita, integrando queste opportunità nella narrazione più ampia della strategia di sviluppo.

Approfondisci le metriche di crescita e i KPI specificati dall'utente, spiegando come questi strumenti di misurazione guideranno l'azienda nel monitoraggio del progresso verso gli obiettivi a lungo termine. Descrivi le considerazioni sulla sostenibilità e responsabilità sociale menzionate come elementi che arricchiscono e completano la strategia complessiva.

Concludi con una riflessione su come l'intera strategia di crescita sia realistica e supportata dalle capacità dell'azienda descritte nel business plan, presentando un futuro promettente ma raggiungibile per l'organizzazione.
        """,
        
        "Finanza": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi in forma narrativa e discorsiva. NON utilizzare elenchi numerati o puntati per il testo descrittivo. Tuttavia, DEVI includere la tabella finanziaria fornita esattamente come specificato nel prompt.

Basandoti sui dati finanziari forniti dall'utente, presenta una sintesi finanziaria professionale e narrativa che dimostri la solidità economica del progetto. Non inventare cifre o proiezioni, utilizza esclusivamente i dati inseriti nei form per costruire una storia finanziaria convincente.

Inizia spiegando il modello di ricavi dell'azienda attraverso una descrizione dettagliata che utilizzi le informazioni fornite sul business model. Racconta come l'azienda genera valore economico, presentando la struttura dei costi principali indicata dall'utente e spiegando come l'investimento richiesto supporterà strategicamente la crescita pianificata.

Presenta quindi la tabella delle proiezioni finanziarie a 5 anni, che rappresenta il cuore dell'analisi economica:

{tabella_finanziaria_md}

Dopo aver mostrato i dati, sviluppa un'analisi narrativa approfondita che commenti sinteticamente ma professionalmente i numeri della tabella. Racconta la storia che emerge dai dati, evidenziando i punti salienti quali la crescita dei ricavi, l'evoluzione dei margini, il raggiungimento del break-even e le proiezioni del cash flow. Presenta questi elementi come capitoli di una storia di successo finanziario, spiegando le dinamiche che sottendono ciascun trend.

Continua con una valutazione della sostenibilità finanziaria del progetto, basandoti esclusivamente sui numeri forniti per costruire un'argomentazione solida e professionale. Spiega attraverso una narrazione convincente perché questi dati rendono l'investimento attraente per il settore bancario, evidenziando gli elementi di sicurezza e le prospettive di crescita che emergono dall'analisi.

Concludi mantenendo la sintesi focalizzata sui dati concreti senza speculazioni, dimostrando attraverso una descrizione professionale e ben argomentata come i numeri supportino chiaramente la fattibilità del business plan e rappresentino un'opportunità solida per gli investitori.
        """,
        
        "Rischio e Mitigazione": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista. Presenta l'analisi dei rischi attraverso paragrafi fluidi e professionali.

Presenta un'analisi professionale e approfondita dei rischi e delle strategie di mitigazione, basandoti esclusivamente sui rischi identificati dall'utente e sulle misure di controllo specificate. Non aggiungere rischi non menzionati, ma elabora professionalmente le informazioni fornite costruendo una valutazione completa e rassicurante del profilo di rischio aziendale.

Inizia utilizzando l'analisi dei rischi fornita dall'utente per presentare attraverso una narrazione professionale i principali fattori di rischio che l'azienda ha identificato. Per ogni categoria di rischio specificata, sviluppa una descrizione dettagliata dell'impatto potenziale e della probabilità di occorrenza, basandoti sulle valutazioni indicate dall'utente e presentando questi elementi come parte di una strategia di risk management matura e consapevole.

Continua presentando le strategie di mitigazione descritte dall'utente per ciascun rischio identificato, spiegando attraverso una narrazione convincente come queste misure riducono concretamente l'esposizione e proteggono la continuità del business. Racconta come l'azienda ha sviluppato questi approcci di mitigazione, mostrando la profondità dell'analisi e la solidità delle soluzioni implementate.

Approfondisci i piani di contingenza specificati dall'utente, descrivendo attraverso una narrazione dettagliata come l'azienda intende rispondere efficacemente a eventuali scenari avversi. Presenta questi piani come dimostrazione della preparazione manageriale e della capacità di adattamento dell'organizzazione di fronte alle sfide.

Sviluppa il tema dei sistemi di monitoraggio e controllo menzionati dall'utente, spiegando come questi permettono di identificare precocemente i segnali di rischio e di intervenire tempestivamente. Descrivi le risorse allocate per la gestione del rischio secondo quanto indicato nei form, presentando questi investimenti come elementi che rafforzano la solidità complessiva del business plan.

Concludi con una valutazione delle informazioni sulle assicurazioni o altre forme di protezione specificate, spiegando come l'azienda trasferisce o condivide strategicamente alcuni rischi per ottimizzare il proprio profilo complessivo. Presenta l'intera strategia di gestione del rischio come un elemento che supporta la sostenibilità a lungo termine dell'azienda, dimostrando come l'approccio proposto sia appropriato e ben calibrato per il profilo di rischio specifico dell'organizzazione.
        """,

        # === SEZIONI PER VERSIONE BREVE/SINTETICA ===

        "Contesto e Mercato": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista.

Presenta un'analisi sintetica ma completa del contesto di mercato in cui opera l'azienda, integrando sia l'analisi della situazione che la strategia di marketing in una narrazione coerente e professionale.

Inizia descrivendo il settore di riferimento e il panorama competitivo, spiegando chiaramente quale opportunità di mercato l'azienda intende cogliere. Utilizza l'analisi SWOT fornita per evidenziare i principali punti di forza che differenziano l'offerta, contestualizzando le opportunità nel mercato attuale.

Prosegui presentando la proposta di valore e il posizionamento strategico dell'azienda. Descrivi i segmenti di clienti target e spiega come la strategia di marketing e i canali di distribuzione permetteranno di raggiungere efficacemente questi mercati. Integra informazioni sul modello di pricing e sulla strategia di acquisizione clienti, mostrando come questi elementi si combinano per creare un approccio al mercato solido e scalabile.

Concludi con una riflessione sulle tendenze di mercato favorevoli e su come l'azienda sia posizionata per capitalizzare queste opportunità, pur rimanendo consapevole delle sfide competitive da affrontare.
        """,

        "Piano Operativo e Management": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista.

Presenta una sintesi integrata del piano operativo, del team di gestione e della strategia di crescita in una narrazione coesa che dimostri la capacità esecutiva dell'azienda.

Inizia descrivendo la struttura operativa essenziale dell'azienda e i processi chiave che garantiscono la delivery del prodotto o servizio. Spiega come l'organizzazione è strutturata per supportare la crescita prevista, evidenziando gli aspetti critici dell'infrastruttura tecnologica e delle risorse necessarie.

Prosegui presentando il team di gestione, focalizzandoti sulle competenze ed esperienze chiave dei fondatori e dei leader principali. Descrivi come la composizione del team e la distribuzione delle responsabilità garantiscano l'esecuzione efficace del business plan. Menziona i piani di espansione del team se rilevanti per la crescita.

Sviluppa quindi la visione di crescita a medio termine, spiegando le principali milestone pianificate e come l'azienda intende scalare le operazioni. Descrivi le opportunità di espansione o diversificazione previste, mostrando come queste si allineino con le capacità operative e le competenze del team.

Concludi evidenziando come l'insieme di operazioni solide, leadership competente e strategia di crescita chiara rendano credibile e realizzabile il piano di business proposto.
        """,

        "Proiezioni Finanziarie": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi in forma narrativa e discorsiva. NON utilizzare elenchi numerati o puntati per il testo descrittivo. Tuttavia, DEVI includere la tabella finanziaria fornita esattamente come specificato.

Presenta una sintesi finanziaria concisa ma completa che dimostri la solidità economica del progetto, basandoti esclusivamente sui dati forniti dall'utente.

Inizia spiegando brevemente il modello di ricavi e la struttura dei costi principali, evidenziando l'investimento richiesto e il suo utilizzo strategico. Presenta quindi la tabella delle proiezioni finanziarie a 5 anni:

{tabella_finanziaria_md}

Dopo la tabella, sviluppa un'analisi narrativa sintetica ma professionale dei dati. Evidenzia la crescita dei ricavi prevista, l'evoluzione dei margini operativi e il punto di break-even. Spiega come le proiezioni dimostrino la sostenibilità economica del progetto e perché rappresentino un'opportunità attraente per gli investitori.

Concludi con una valutazione concisa della solidità finanziaria del business plan, basandoti sui numeri concreti per costruire un'argomentazione convincente sull'opportunità di investimento.
        """,

        "Rischi Principali": """
Agisci come un consulente finanziario esperto di 20 anni.

IMPORTANTE: Scrivi ESCLUSIVAMENTE in forma narrativa e discorsiva. NON utilizzare mai elenchi numerati (1, 2, 3), elenchi puntati (-, •), o qualsiasi tipo di lista.

Presenta un'analisi sintetica ma professionale dei principali rischi identificati e delle strategie di mitigazione adottate dall'azienda.

Inizia identificando i rischi più significativi che potrebbero impattare il business, basandoti sulle informazioni fornite dall'utente. Descrivi brevemente la natura di questi rischi e il loro potenziale impatto sull'azienda, mantenendo un tono professionale e bilanciato.

Prosegui presentando le principali strategie di mitigazione implementate per ciascuna categoria di rischio rilevante. Spiega come l'azienda intende ridurre l'esposizione attraverso misure preventive concrete e piani di contingenza appropriati. Evidenzia i sistemi di monitoraggio che permetteranno di identificare precocemente eventuali segnali di allerta.

Concludi rassicurando sulla maturità dell'approccio alla gestione del rischio, mostrando come la consapevolezza e la preparazione dell'azienda rafforzino la solidità complessiva del business plan e la fiducia degli investitori nella capacità di navigare le sfide future.
        """
    }

    prompt = istruzioni_specifiche.get(nome_sezione, "...")

    # Aggiungi prefisso per versione breve
    if tipo_bp == 'breve':
        prompt = prefisso_breve + "\n" + prompt

    # Gestione tabelle finanziarie per entrambe le versioni
    if nome_sezione == "Finanza" or nome_sezione == "Proiezioni Finanziarie":
        print(f">>>> PRIMA sostituzione: {tabella_finanziaria_md[:100]}...")  # Debug: mostra i primi 100 caratteri della tabella
        # Inserisci la tabella finanziaria nel prompt
        prompt = prompt.replace("{tabella_finanziaria_md}", tabella_finanziaria_md)
        print(f">>>> DOPO sostituzione: tabella presente nel prompt = {'tabella_finanziaria_md' not in prompt}")

    return prompt
    
    
def costruisci_prompt_per_sezione(nome_sezione, dati_bp, contesto_precedente, tabella_finanziaria_md="", tipo_bp="completo"):
    prompt_parts = []

    # 1. Ottieni le istruzioni specifiche per la sezione corrente
    istruzione_specifica = ottieni_istruzioni_per_sezione(nome_sezione, tabella_finanziaria_md, tipo_bp)
    prompt_parts.append(istruzione_specifica)

    # 2. Fornisci il contesto delle sezioni precedenti (fondamentale per la coerenza)
    if contesto_precedente:
        prompt_parts.append("\n--- CONTESTO (SEZIONI GIÀ SCRITTE, da usare per coerenza) ---")
        prompt_parts.append(contesto_precedente)
        prompt_parts.append("--- FINE CONTESTO ---")

    # 3. Fornisci i dati grezzi dell'utente come riferimento generale
    dati_utente_formattati = []
    for key, value in dati_bp.items():
        dati_utente_formattati.append(f"[{key.upper()}]: {value}")
    prompt_parts.append("\n--- DATI UTENTE COMPLETI (da usare come riferimento) ---")
    prompt_parts.append("\n".join(dati_utente_formattati))
    prompt_parts.append("--- FINE DATI UTENTE ---")

    return "\n".join(prompt_parts)


# Esempio d'uso (FUORI DA TUTTE LE CLASSI)
if __name__ == "__main__":
    # Test della funzione
    prompt = ottieni_istruzioni_per_sezione("Marketing")
    print(prompt)
    
    # Test con tabella finanziaria
    tabella = "|Anno|Ricavi|Costi|\n|2024|100K|80K|"
    prompt_finanza = ottieni_istruzioni_per_sezione("Finanza", tabella)
    print(prompt_finanza)

@app.route('/genera-business-plan', methods=['POST'])
def handle_genera_business_plan():
    # 1. Controlla e ricevi i dati JSON dal frontend
    if not request.is_json:
        return jsonify({"status": "error", "message": "Richiesta non in formato JSON"}), 400
    
    dati_completi_bp = request.get_json()
    if not dati_completi_bp:
        return jsonify({"status": "error", "message": "Nessun dato ricevuto"}), 400
    
    if not anthropic_client:
        return jsonify({"status": "error", "message": "Client Anthropic non inizializzato."}), 500

    print(">>> Avvio generazione business plan per sezioni...")

    # 2. Inizia il blocco try per gestire qualsiasi errore durante il processo
    try:
        # 3. Determina il tipo di business plan richiesto
        tipo_bp = dati_completi_bp.get('tipoBP', 'completo')
        print(f">>> Tipo di Business Plan richiesto: {tipo_bp}")

        # 4. Definisci le sezioni in base al tipo di BP
        if tipo_bp == 'breve':
            sezioni_da_generare = [
                'Riassunto Esecutivo',
                'Contesto e Mercato',
                'Piano Operativo e Management',
                'Proiezioni Finanziarie',
                'Rischi Principali'
            ]
            max_tokens = 2500  # Versione più breve
            print(">>> Modalità SINTETICA attivata: 5 sezioni, max 5 pagine")
        else:
            sezioni_da_generare = [
                'Riassunto Esecutivo',
                'Analisi della Situazione',
                'Marketing',
                'Operazioni',
                'Gestione',
                'Strategia di Crescita',
                'Finanza',
                'Rischio e Mitigazione'
            ]
            max_tokens = 4000  # Versione completa
            print(">>> Modalità COMPLETA attivata: 8 sezioni dettagliate")

        business_plan_completo = []
        contesto_precedente = ""

        # Pre-calcola la tabella finanziaria una sola volta
        dati_finanziari = dati_completi_bp.get('datiFinanziari', {})
        valuta_scelta = dati_completi_bp.get('datiInvestimenti', {}).get('valuta', 'EUR')
        valuta_simbolo = valuta_scelta.split('(')[1].replace(')','') if '(' in valuta_scelta else valuta_scelta
        tabella_markdown = calcola_e_formatta_proiezioni(dati_finanziari, valuta_simbolo)

        # 5. Esegui il ciclo per generare ogni sezione
        for i, nome_sezione in enumerate(sezioni_da_generare):
            print(f">>> Generazione Sezione {i+1}/{len(sezioni_da_generare)}: '{nome_sezione}'...")

            # Costruisci il prompt specifico per la sezione corrente
            prompt_da_usare = costruisci_prompt_per_sezione(nome_sezione, dati_completi_bp, contesto_precedente, tabella_markdown, tipo_bp)

            # Chiamata API a Claude per la sezione corrente
            # Assicurati di usare il tuo model_id corretto
            model_id = "claude-3-7-sonnet-20250219"  # Sostituisci con il tuo model_id se necessario
            response = anthropic_client.messages.create(
                model=model_id,
                max_tokens=max_tokens,
                messages=[{"role": "user", "content": prompt_da_usare}]
            )
            testo_sezione_generata = response.content[0].text.strip() if response.content else ""
            
                       
            sezione_formattata = f"## {nome_sezione}\n\n{testo_sezione_generata}\n\n"
            business_plan_completo.append(sezione_formattata)
            contesto_precedente += sezione_formattata
            
            print(f">>> Sezione '{nome_sezione}' generata con successo.")

        # 5. Prepara il risultato finale
        testo_finale_completo = "".join(business_plan_completo)
        html_output = convert_text_to_html(testo_finale_completo)

        return jsonify({
            "status": "success",
            "business_plan_text": testo_finale_completo,
            "business_plan_html": html_output
        })

    # 6. Blocco except per catturare qualsiasi errore avvenuto nel blocco try
    except Exception as e:
        print(f"ERRORE CRITICO durante la generazione a sezioni: {e}")
        # Includi un traceback per un debug più facile nel terminale
        import traceback
        traceback.print_exc()
        return jsonify({"status": "error", "message": f"Errore interno del server durante la generazione del business plan: {e}"}), 500




# Rotta per la pagina iniziale (landing page)
@app.route('/')
def index():
    # Assumendo che il tuo file principale si chiami index.html e sia in templates/
    return render_template('index.html')

# === ROUTE DASHBOARD (PROTETTA) ===
@app.route('/dashboard')
@login_required
def dashboard():
    """
    Dashboard utente - mostra crediti, business plan generati, storico ordini
    """
    # Ottieni i crediti dell'utente
    crediti_brevi = current_user.get_credits('breve')
    crediti_completi = current_user.get_credits('completo')

    # Ottieni i business plan generati dall'utente
    from models import BusinessPlan
    business_plans = BusinessPlan.query.filter_by(user_id=current_user.id).order_by(BusinessPlan.created_at.desc()).all()

    # Ottieni gli ordini
    from models import Order
    ordini = Order.query.filter_by(user_id=current_user.id).order_by(Order.created_at.desc()).all()

    return render_template('dashboard.html',
                         crediti_brevi=crediti_brevi,
                         crediti_completi=crediti_completi,
                         business_plans=business_plans,
                         ordini=ordini,
                         user=current_user)


# === ROUTE ACQUISTO CREDITI ===
@app.route('/buy-credits', methods=['GET', 'POST'])
@login_required
def buy_credits():
    """
    Pagina acquisto crediti e gestione ordine
    """
    if request.method == 'GET':
        # Mostra la pagina con i pacchetti disponibili
        pacchetti = get_pacchetti_crediti()
        crediti_brevi = current_user.get_credits('breve')
        crediti_completi = current_user.get_credits('completo')

        return render_template('buy_credits.html',
                             pacchetti=pacchetti,
                             crediti_brevi=crediti_brevi,
                             crediti_completi=crediti_completi)

    # POST - Processa l'acquisto
    pacchetto_id = request.form.get('pacchetto_id')

    if not pacchetto_id:
        flash('Seleziona un pacchetto', 'error')
        return redirect(url_for('buy_credits'))

    # Ottieni informazioni pacchetto
    pacchetto = get_pacchetto_by_id(pacchetto_id)

    if not pacchetto:
        flash('Pacchetto non trovato', 'error')
        return redirect(url_for('buy_credits'))

    # IMPORTANTE: Per ora saltiamo Stripe e aggiungiamo crediti direttamente
    # In STEP 4 implementeremo Stripe per pagamenti reali
    try:
        # Crea ordine (stato: completato direttamente per test)
        ordine = Order(
            user_id=current_user.id,
            pacchetto_id=pacchetto_id,
            tipo_credito=pacchetto['tipo'],
            quantita=pacchetto['quantita'],
            prezzo=pacchetto['prezzo'],
            stato='completato',  # Per ora sempre completato
            metodo_pagamento='test'  # Segnaposto per test
        )
        db.session.add(ordine)
        db.session.commit()

        # Aggiungi i crediti all'utente
        aggiungi_crediti(
            user_id=current_user.id,
            tipo=pacchetto['tipo'],
            quantita=pacchetto['quantita'],
            order_id=ordine.id
        )

        flash(f"✅ Acquisto completato! Aggiunti {pacchetto['quantita']} crediti '{pacchetto['tipo']}'", 'success')
        return redirect(url_for('dashboard'))

    except Exception as e:
        db.session.rollback()
        flash(f'Errore durante l\'acquisto: {str(e)}', 'error')
        return redirect(url_for('buy_credits'))


# === NUOVE ROTTE PER I FORM ===
# Aggiungi una rotta per ogni pagina del tuo form.
# Flask cercherà i file partendo dalla cartella 'templates'.
# Visto che hai messo i form in 'templates/forms/', il percorso sarà 'forms/nomefile.html'.

@app.route('/forms/<page_name>')
def mostra_form(page_name):
    # Rotta dinamica per servire tutti i file dalla cartella forms
    return render_template(f'forms/{page_name}.html')


@app.route('/scarica-pdf', methods=['POST'])
def scarica_pdf():
    print(">>> 1. Richiesta a /scarica-pdf ricevuta.")
    
    try:
        data = request.get_json()
        raw_text = data.get('text')
        if not raw_text:
            return jsonify({"error": "Nessun testo ricevuto"}), 400

        print(">>> 2. Testo ricevuto correttamente.")

        html_content = convert_text_to_html(raw_text)
        print(">>> 3. Conversione a HTML completata.")

        full_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="UTF-8">
            <title>Business Plan</title>
        </head>
        <body>
            <h1 class="main-title">Business Plan</h1>
            {html_content}
        </body>
        </html>
        """
        print(">>> 4. Stringa HTML completa creata.")

        # USA IL FILE CSS CORRETTO!
        print(">>> 5. Sto per caricare il CSS con WeasyPrint...")
        css_file = CSS(filename='static/pdf_style.css')  # <-- CORRETTO!
        print(">>> 6. Caricamento CSS RIUSCITO.")

        print(">>> 7. Sto per generare il PDF con WeasyPrint...")
        pdf_bytes = HTML(string=full_html).write_pdf(stylesheets=[css_file])
        print(">>> 8. Generazione PDF RIUSCITA.")

        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={"Content-Disposition": "attachment;filename=business_plan.pdf"}
        )

    except Exception as e:
        print(f"!!! ERRORE PYTHON CATTURATO: {e}")
        return jsonify({"error": f"Errore nella generazione del PDF: {str(e)}"}), 500
# Sostituisci l'intera rotta @app.route('/scarica-docx', ...) con questa:

@app.route('/scarica-docx', methods=['POST'])
def scarica_docx():
    """
    Endpoint che converte Markdown in HTML, e poi da HTML a DOCX per coerenza.
    """
    try:
        data = request.get_json()
        raw_text = data.get('text')
        if not raw_text:
            return jsonify({"error": "Nessun testo ricevuto"}), 400

        # 1. Converti il Markdown in un HTML pulito e standard
        html_content = convert_text_to_html(raw_text)

        # 2. Usa BeautifulSoup per "leggere" l'HTML
        soup = BeautifulSoup(html_content, 'html.parser')
        
        document = Document()

        # 3. Itera sugli elementi HTML e costruisci il documento Word
        for element in soup.find_all(['h1', 'h2', 'h3', 'p', 'ul', 'ol', 'table']):
            if element.name == 'h1':
                document.add_heading(element.text, level=1)
            elif element.name == 'h2':
                document.add_heading(element.text, level=2)
            elif element.name == 'h3':
                document.add_heading(element.text, level=3)
            elif element.name == 'p':
                # Gestisce paragrafi con testo normale, <strong> e <em>
                p = document.add_paragraph()
                for child in element.children:
                    if child.name == 'strong':
                        p.add_run(child.text).bold = True
                    elif child.name == 'em':
                        p.add_run(child.text).italic = True
                    else:
                        p.add_run(str(child))
            elif element.name == 'ul':
                for li in element.find_all('li'):
                    document.add_paragraph(li.text, style='List Bullet')
            elif element.name == 'ol':
                for li in element.find_all('li'):
                    document.add_paragraph(li.text, style='List Number')
            elif element.name == 'table':
                # Aggiunge una tabella al documento
                table_data = []
                for row in element.find_all('tr'):
                    row_data = [cell.text for cell in row.find_all(['th', 'td'])]
                    table_data.append(row_data)
                
                if table_data:
                    table = document.add_table(rows=len(table_data), cols=len(table_data[0]))
                    table.style = 'Table Grid'
                    for i, row_data in enumerate(table_data):
                        for j, cell_text in enumerate(row_data):
                            table.cell(i, j).text = cell_text

        # 4. Salva e invia il file
        memoria_file = io.BytesIO()
        document.save(memoria_file)
        memoria_file.seek(0)
        return send_file(memoria_file, as_attachment=True, download_name='business_plan.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

    except Exception as e:
        print(f"Errore durante la generazione del DOCX: {e}")
        return jsonify({"error": "Impossibile generare il DOCX"}), 500



if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)